# SPDX-License-Identifier: EUPL-1.2
"""Werkboeken lezen: de jaartotalen met drie kopregels en de detailtabbladen per meetpunt.

De lezers geven de bron **getrouw** terug (bronkolomnamen, bronwaarden); vertalen, typeren naar
``physiquant__unit``-namen en vlaggen gebeurt in ``tabellen.py`` en ``toetsen.py``. Werkboeken worden
met ``openpyxl`` in ``read_only``-modus en ``data_only=True`` gelezen: formules (het tabblad
``summary``) komen niet mee, alleen opgeslagen waarden, en het geheugen blijft beperkt.
"""

from __future__ import annotations

import logging
import warnings
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterator

import numpy as np
import openpyxl
import pandas as pd

log = logging.getLogger(__name__)

#: Het rekenblad van Wattopia; geen meetdata (ontwerp §1.2 en §8).
REKENBLAD = "summary"
TIJDKOLOM_BRON = "Timestamp"


@dataclass
class Jaartotalen:
    """De drie tabbladen van het werkboek met jaartotalen."""

    projectdata: pd.DataFrame
    monitoringdata: pd.DataFrame
    #: toegestane waarden per keuzekolom van het tabblad ``keuzemenu``
    keuzemenu: dict[str, list[str]]
    #: per tabblad: bronkolomnaam naar groepsnaam (kopregel 1) en naar eenheid of type (kopregel 3)
    groepen: dict[str, dict[str, str]] = field(default_factory=dict)
    eenheden: dict[str, dict[str, str]] = field(default_factory=dict)


def _open(pad: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)  # "Data Validation extension is not supported"
        return openpyxl.load_workbook(pad, read_only=True, data_only=True)


def _is_leeg(waarde) -> bool:
    return waarde is None or (isinstance(waarde, str) and waarde.strip() == "")


def lees_gegroepeerd_blad(rijen: list[tuple], sleutelkolom: str | None = None
                          ) -> tuple[pd.DataFrame, dict[str, str], dict[str, str]]:
    """Een blad met drie kopregels (groep, kolomnaam, eenheid of type) en data vanaf rij 4.

    Lege cellen blijven leeg (``None``/``NaN``): "niet gerapporteerd", geen nul. Volledig lege
    rijen (en, met ``sleutelkolom``, rijen zonder sleutel) worden overgeslagen.
    """
    if len(rijen) < 3:
        raise ValueError("Verwacht minstens drie kopregels")
    groeprij, naamrij, eenheidrij = rijen[0], rijen[1], rijen[2]
    aantal = max((i + 1 for i, n in enumerate(naamrij) if not _is_leeg(n)), default=0)
    namen = [str(n).strip() for n in naamrij[:aantal]]

    groepen: dict[str, str] = {}
    huidige = ""
    for i, naam in enumerate(namen):
        g = groeprij[i] if i < len(groeprij) else None
        if not _is_leeg(g):
            huidige = str(g).strip()
        groepen[naam] = huidige
    eenheden = {
        naam: ("" if i >= len(eenheidrij) or _is_leeg(eenheidrij[i]) else str(eenheidrij[i]).strip())
        for i, naam in enumerate(namen)
    }

    data = []
    for rij in rijen[3:]:
        rij = tuple(rij[:aantal]) + (None,) * max(0, aantal - len(rij))
        if all(_is_leeg(c) for c in rij):
            continue
        if sleutelkolom is not None and _is_leeg(rij[namen.index(sleutelkolom)]):
            continue
        data.append(rij)
    df = pd.DataFrame(data, columns=namen, dtype=object)
    # Kolommen die alleen getallen (of leeg) bevatten worden numeriek; tekst blijft tekst.
    for kolom in df.columns:
        waarden = df[kolom]
        gevuld = waarden.dropna()
        if len(gevuld) and all(isinstance(w, (int, float, np.integer, np.floating))
                               and not isinstance(w, bool) for w in gevuld):
            df[kolom] = pd.to_numeric(waarden)
        elif len(gevuld) == 0:
            df[kolom] = pd.Series([None] * len(df), dtype=object)
        else:
            df[kolom] = waarden.map(lambda w: None if _is_leeg(w) else str(w))
    return df, groepen, eenheden


def lees_jaartotalen(pad: Path) -> Jaartotalen:
    """Lees ``Projectdata``, ``Monitoringdata`` en ``keuzemenu`` uit het werkboek met jaartotalen."""
    wb = _open(Path(pad))
    try:
        blad = {}
        for naam in ("Projectdata", "Monitoringdata", "keuzemenu"):
            if naam not in wb.sheetnames:
                raise ValueError(f"Tabblad {naam!r} ontbreekt in {Path(pad).name}")
            blad[naam] = list(wb[naam].iter_rows(values_only=True))
    finally:
        wb.close()

    projecten, g1, e1 = lees_gegroepeerd_blad(blad["Projectdata"], sleutelkolom="Projectcode")
    woningen, g2, e2 = lees_gegroepeerd_blad(blad["Monitoringdata"], sleutelkolom="WONINGCODE")

    menu_rijen = blad["keuzemenu"]
    keuzemenu: dict[str, list[str]] = {}
    if menu_rijen:
        kop = menu_rijen[0]
        for i, naam in enumerate(kop):
            if _is_leeg(naam):
                continue
            keuzemenu[str(naam).strip()] = [
                str(r[i]) for r in menu_rijen[1:] if i < len(r) and not _is_leeg(r[i])
            ]
    return Jaartotalen(
        projectdata=projecten, monitoringdata=woningen, keuzemenu=keuzemenu,
        groepen={"Projectdata": g1, "Monitoringdata": g2},
        eenheden={"Projectdata": e1, "Monitoringdata": e2},
    )


def _typeer_detailkolom(waarden: pd.Series) -> pd.Series:
    """Getallen blijven getallen (ontbrekend wordt NaN); een kolom met tekst wordt tekst (``None``
    voor ontbrekend). Niets wordt afgerond of vervangen."""
    gevuld = waarden.dropna()
    if len(gevuld) == 0:
        return pd.Series(np.full(len(waarden), np.nan), index=waarden.index, dtype="float64")
    if all(isinstance(w, (int, float, np.integer, np.floating, bool)) for w in gevuld):
        if len(gevuld) == len(waarden) and all(isinstance(w, (int, np.integer, bool)) for w in gevuld):
            return waarden.astype("int64")
        return pd.to_numeric(waarden).astype("float64")
    return waarden.map(lambda w: None if w is None or (isinstance(w, float) and np.isnan(w)) else str(w))


def lees_detailtabblad(rijen: Iterator[tuple]) -> pd.DataFrame | None:
    """Een tabblad met een kopregel en daaronder de metingen; ``None`` als er geen metingen zijn."""
    try:
        kop = next(rijen)
    except StopIteration:
        return None
    aantal = max((i + 1 for i, n in enumerate(kop) if not _is_leeg(n)), default=0)
    if aantal == 0:
        return None
    namen = [str(n).strip() for n in kop[:aantal]]
    data = []
    for rij in rijen:
        rij = tuple(rij[:aantal]) + (None,) * max(0, aantal - len(rij))
        if all(_is_leeg(c) for c in rij):
            continue
        data.append(rij)
    if not data:
        return None
    df = pd.DataFrame(data, columns=namen, dtype=object)
    for kolom in df.columns:
        df[kolom] = _typeer_detailkolom(df[kolom])
    return df


def lees_detailwerkboek(pad: Path) -> Iterator[tuple[str, pd.DataFrame | None]]:
    """Geef per meetpunt ``(tabbladnaam, DataFrame)``; ``None`` voor een tabblad zonder metingen.

    Het rekenblad ``summary`` wordt overgeslagen. Het werkboek blijft open tot de generator klaar is,
    dus consumeer hem in één keer en schrijf elk tabblad direct weg; zo staat er maar één tabblad
    tegelijk in het geheugen.
    """
    wb = _open(Path(pad))
    try:
        for naam in wb.sheetnames:
            if naam == REKENBLAD:
                continue
            df = lees_detailtabblad(wb[naam].iter_rows(values_only=True))
            if df is not None and TIJDKOLOM_BRON not in df.columns:
                raise ValueError(f"Tabblad {naam!r} in {Path(pad).name} heeft geen kolom {TIJDKOLOM_BRON!r}")
            yield naam, df
    finally:
        wb.close()
