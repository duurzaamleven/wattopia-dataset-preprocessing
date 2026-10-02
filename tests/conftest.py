"""Gedeelde hulpmiddelen: een klein synthetisch pakket (nooit de echte data) in de vorm van de bron.

Het pakket heeft twee woningen met tijdreeksen van twee dagen, met metingen op veelvouden van 300 s,
zodat de som over het raster precies gelijk is aan de eindstand min de beginstand.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

from wattopia_preprocessing import kolommen as K
from wattopia_preprocessing import tijd as T

AANTAL = 576  # intervallen van vijf minuten in twee dagen
START = pd.Timestamp("2019-03-01 00:00", tz=K.TIJDZONE)
START_S = int(START.tz_convert("UTC").timestamp())

#: Vermogen per vijf minuten in de tellers, per woning: factor op de basisstappen hieronder.
WONINGEN = {"W_aaaaaaaa": 1.0, "W_bbbbbbbb": 1.5}
#: Stap per vijf minuten in de teller.
STAPPEN = {"hp": 0.1, "booster": 0.01, "pv": 0.2, "wtw": 0.005, "q": 0.001}


def verwachte_totalen(factor: float) -> dict[str, float]:
    """De totalen over de 576 intervallen, in de eenheid van het jaartotaal."""
    return {
        "e_opwekker_totaal": (STAPPEN["hp"] + STAPPEN["booster"]) * AANTAL * factor,
        "q_ruimte": STAPPEN["q"] * AANTAL * factor,
        "e_opwek": STAPPEN["pv"] * AANTAL * factor,
        "e_gebouw_overig": STAPPEN["wtw"] * AANTAL * factor,
    }


def kort_raster() -> pd.DatetimeIndex:
    return pd.date_range(START, periods=AANTAL, freq="300s")


def _tijden(n: int = AANTAL + 1) -> np.ndarray:
    return START_S + 300 * np.arange(n)


def _teller_blad(ws, tijden, stap, kolom="total_energy_in", ook_vermogen=True):
    kop = ["Timestamp"] + (["power"] if ook_vermogen else []) + [kolom]
    ws.append(kop)
    for i, t in enumerate(tijden):
        ws.append([int(t)] + ([stap * 12000] if ook_vermogen else []) + [stap * i])


def schrijf_detailwerkboek(pad: Path, factor: float) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["Totaal", "=SUM(A2:A3)"])
    tijden = _tijden()

    ws = wb.create_sheet("HeatPump")
    ws.append(["Timestamp", "booster_status", "error_code", "firmware_version", "flow_temp", "op_mode",
               "room_temp", "temp"])
    for i, t in enumerate(tijden):
        flow = -100.0 if i == 10 else 35.0 + (i % 5)
        room = 0.0 if i == 20 else 21.0
        ws.append([int(t), i % 2, "E" + str(i % 3) if i % 3 else None, "1.2.3", flow, 3, room, 40.0])

    for naam, sleutel in (("energyHeatpump", "hp"), ("energyBooster", "booster"), ("energyImmersion", None),
                          ("energyWtwReg", "wtw")):
        ws = wb.create_sheet(naam)
        _teller_blad(ws, tijden, 0.0 if sleutel is None else STAPPEN[sleutel] * factor)

    ws = wb.create_sheet("solar")
    ws.append(["Timestamp", "power", "total_energy_in", "total_energy_out"])
    for i, t in enumerate(tijden):
        ws.append([int(t), -STAPPEN["pv"] * factor * 12000, 0.0, STAPPEN["pv"] * factor * i])

    ws = wb.create_sheet("flowHeatSpaceHeating")
    ws.append(["Timestamp", "flow_temp", "return_temp", "total_energyGJ", "volume_out"])
    for i, t in enumerate(tijden):
        ws.append([int(t), 30.0, 25.0, STAPPEN["q"] * factor * i, i * 0.001])

    ws = wb.create_sheet("sheet15")
    ws.append(["Timestamp"])  # leeg tabblad: alleen een kopregel
    wb.save(pad)


def schrijf_jaartotalen(pad: Path, afwijking_pct: float = 0.0) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Projectdata"
    namen = [k.bron for k in K.PROJECTEN]
    ws.append(["PROJECTIDENTIFICATIE"] + [None] * 2 + ["ALGEMENE GEGEVENS"] + [None] * (len(namen) - 4))
    ws.append(namen)
    ws.append([None, "aantal"] + [None] * (len(namen) - 2))
    rij = {n: None for n in namen}
    rij.update({"Projectcode": "PROJ000001", "Monitoringdata": 2, "KNMI Weerstation": "De Bilt",
                "Bouwjaar afgerond": 1970, "Renovatiejaar afgerond": 2019, "Gebruiksoppervlak afgerond": 100,
                "Woningtypologie": "Grondgebonden jaren ’60", "EPV": "ja", "Actieve koeling": "nee",
                "Warmteopwekker": "Warmtepomp", "Afgiftesysteem": "Een afgiftesysteem dat niet bestaat"})
    ws.append([rij[n] for n in namen])

    ws = wb.create_sheet("Monitoringdata")
    ws.append(["WONINGIDENTIFICATIE", None, None, "ENERGIEPRESTATIES"] + [None] * 18)
    namen = [k.bron for k in K.WONINGEN]
    ws.append(namen)
    eenheden = {"kWh_a_1": "kWh_e / jaar", "GJ_a_1": "GJ / jaar", "ppm": "ppm", "pct": "percentage",
                "m3_a_1": "m3 / jaar", "degC": "graden C"}
    ws.append([None, None, "jaartal"] + [eenheden.get(K.stam_en_eenheid(k.naam)[1]) for k in K.WONINGEN[3:]])
    for code, factor in WONINGEN.items():
        t = verwachte_totalen(factor)
        waarden = {k.bron: None for k in K.WONINGEN}
        waarden.update({"PROJECTCODE": "PROJ000001", "WONINGCODE": code, "JAARTAL MONITORING": 2019,
                        "Opgewekte energie": t["e_opwek"] * (1 + afwijking_pct / 100),
                        "Totaal gebruikte energie door de warmteopwekker(s)": t["e_opwekker_totaal"],
                        "Geleverde warmte voor ruimteverwarming": t["q_ruimte"],
                        "Gebruikte energie voor overige gebouwgebonden functies": t["e_gebouw_overig"],
                        "Beschikbaar voor persoonlijk gebruik": -50.0,
                        "Setpoint temperatuur woonkamer": 0})
        ws.append([waarden[n] for n in namen])
    waarden = {k.bron: None for k in K.WONINGEN}
    waarden.update({"PROJECTCODE": "PROJ000001", "WONINGCODE": "W_zonderdetail", "JAARTAL MONITORING": 2020,
                    "Opgewekte energie": 1000.0})
    ws.append([waarden[n] for n in namen])

    ws = wb.create_sheet("keuzemenu")
    menu = ["Woningtypologie", "Ja / nee vragen"] + sorted({k.menu for k in K.PROJECTEN if k.menu} - {"Woningtypologie"})
    ws.append(menu)
    ws.append(["Grondgebonden jaren ‘60", "ja"] + ["Warmtepomp" if m == "Warmteopwekker" else "Radiator" if m == "Afgiftesysteem" else "x" for m in menu[2:]])
    ws.append(["Grondgebonden jaren 50", "nee"] + ["y"] * (len(menu) - 2))
    wb.save(pad)


@pytest.fixture(scope="session")
def pakket(tmp_path_factory) -> Path:
    """Een synthetisch ``data/raw`` met het jaartotalenwerkboek en twee detailwerkboeken. Eén keer per
    sessie gebouwd, want het wegschrijven van werkboeken is traag; toetsen die het pakket wijzigen,
    werken op een kopie."""
    raw = tmp_path_factory.mktemp("pakket") / "data" / "raw"
    raw.mkdir(parents=True)
    schrijf_jaartotalen(raw / "Woningprestaties 2022-07-22 ANON.xlsx")
    for code, factor in WONINGEN.items():
        schrijf_detailwerkboek(raw / f"Detaildata woning_{code.removeprefix('W_')}.xlsx", factor)
    return raw
